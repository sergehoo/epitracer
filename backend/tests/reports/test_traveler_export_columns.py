"""Colonnes provenance / destination / séjour des exports voyageurs (CSV, Excel, PDF, hebdo)."""

import csv
import io
from datetime import date, datetime, timedelta

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, RoleAssignment, RoleCode
from apps.geo.models import Country
from apps.quarantine.models import QuarantineRecord
from apps.reports.traveler_fields import (
    ITINERARY_HEADERS,
    NOT_PROVIDED,
    itinerary_row,
    stay_duration_days,
)
from apps.travelers.models import Traveler, TravelHistoryEntry

pytestmark = pytest.mark.django_db

ARRIVAL = date(2026, 10, 1)


@pytest.fixture
def full_traveler(traveler):
    """Voyageur avec provenance, deux séjours récents, destination et séjour en CI daté."""
    traveler.arrival_date = ARRIVAL
    traveler.confinement_city = "Bouaké"
    traveler.save()
    ng, gh, bf = (Country.objects.get(code=c) for c in ("NG", "GH", "BF"))
    TravelHistoryEntry.objects.create(
        traveler=traveler,
        role="origin",
        country=ng,
        city="Lagos",
        province="Lagos State",
    )
    TravelHistoryEntry.objects.create(
        traveler=traveler,
        role="visited",
        country=gh,
        arrival_date=date(2026, 9, 20),
        departure_date=date(2026, 9, 25),
    )
    TravelHistoryEntry.objects.create(
        traveler=traveler,
        role="visited",
        country=bf,
        arrival_date=date(2026, 9, 26),
        departure_date=date(2026, 9, 30),
    )
    TravelHistoryEntry.objects.create(traveler=traveler, role="transit", country=gh)  # ignoré
    # Séjour en Côte d'Ivoire : date de départ connue
    TravelHistoryEntry.objects.create(
        traveler=traveler,
        role="visited",
        country=traveler.entry_point.country,
        departure_date=ARRIVAL + timedelta(days=12),
    )
    return traveler


# ------------------------------------------------------------------ durée


@pytest.mark.parametrize(
    "arrival,departure,expected",
    [
        (date(2026, 10, 1), date(2026, 10, 13), 12),
        (date(2026, 10, 1), date(2026, 10, 1), 0),  # même jour
        (date(2026, 10, 1), None, None),  # départ inconnu
        (None, date(2026, 10, 13), None),  # arrivée inconnue
        (date(2026, 10, 13), date(2026, 10, 1), None),  # incohérent : jamais négatif
    ],
)
def test_stay_duration_days(arrival, departure, expected):
    assert stay_duration_days(arrival, departure) == expected


# ---------------------------------------------------------------- valeurs


def test_columns_with_complete_data(full_traveler):
    values = dict(zip(ITINERARY_HEADERS, itinerary_row(full_traveler), strict=True))
    assert values == {
        "Date d'enregistrement": timezone.localtime(full_traveler.created_at).strftime("%d/%m/%Y"),
        "Pays de provenance": Country.objects.get(code="NG").name,
        # Le plus récent d'abord ; transit et pays de destination exclus
        "Pays de séjour récent": ", ".join(
            [
                Country.objects.get(code="BF").name,
                Country.objects.get(code="GH").name,
            ]
        ),
        "Ville ou province de provenance": "Lagos, Lagos State",
        "Pays de destination": full_traveler.entry_point.country.name,
        "Ville de destination": "Bouaké",
        "Durée du séjour en Côte d'Ivoire (jours)": 12,
    }


def test_missing_information_is_never_invented(entry_point):
    bare = Traveler.objects.create(first_name="Sans", last_name="ITINERAIRE", entry_point=None)
    row = itinerary_row(bare)
    assert row[0] == timezone.localtime(bare.created_at).strftime("%d/%m/%Y")
    assert row[1:] == [NOT_PROVIDED] * 6


def test_duration_requires_both_dates(traveler):
    traveler.arrival_date = None
    traveler.save()
    TravelHistoryEntry.objects.create(
        traveler=traveler,
        role="visited",
        country=traveler.entry_point.country,
        departure_date=ARRIVAL,
    )
    assert itinerary_row(traveler)[-1] == NOT_PROVIDED


def test_free_text_formula_is_neutralized(traveler):
    traveler.confinement_city = '=HYPERLINK("http://x")'
    traveler.save()
    assert itinerary_row(traveler)[5].startswith("'=")


# ------------------------------------------------------------- exports API


@pytest.fixture
def admin_client(superadmin):
    c = APIClient()
    c.force_authenticate(user=superadmin)
    return c


def _query(fmt: str) -> str:
    return f"/api/v1/reports/travelers/?output={fmt}&date_from=2026-09-01&date_to=2026-10-31"


def _csv_rows(resp) -> list[list[str]]:
    text = resp.content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text), delimiter=";"))


def test_csv_and_excel_carry_identical_values(admin_client, full_traveler):
    full_traveler.age = 57
    full_traveler.save(update_fields=["age"])
    csv_resp = admin_client.get(_query("csv"))
    assert csv_resp.status_code == 200
    header, row = _csv_rows(csv_resp)
    # Colonnes existantes conservées, nouvelles colonnes à la fin
    assert header[:2] == ["ID public", "Nom"] and header[-7:] == ITINERARY_HEADERS
    assert row[-7:] == [str(v) for v in itinerary_row(full_traveler)]
    assert row[-1] == "12" and row[-2] == "Bouaké"
    assert row[header.index("Âge")] == "57 ans"  # libellé français, pas « years »

    xlsx_resp = admin_client.get(_query("xlsx"))
    assert xlsx_resp.status_code == 200
    assert xlsx_resp["Content-Disposition"].endswith('.xlsx"')
    ws = load_workbook(io.BytesIO(xlsx_resp.content))["Données"]
    xlsx_header = [c.value for c in ws[1]]
    xlsx_row = [c.value for c in ws[2]]
    assert xlsx_header == header
    assert [str(v) for v in xlsx_row[-7:]] == row[-7:]


def test_pdf_export_renders_wide_table(admin_client, full_traveler):
    resp = admin_client.get(_query("pdf"))
    assert resp.status_code == 200
    assert resp["Content-Type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")


def test_filters_and_missing_values_in_export(admin_client, full_traveler, entry_point):
    other = Traveler.objects.create(first_name="B", last_name="BARE", arrival_date=ARRIVAL)
    rows = _csv_rows(admin_client.get(_query("csv") + f"&entry_point={entry_point.id}"))
    ids = [r[0] for r in rows[1:]]
    assert full_traveler.public_id in ids and other.public_id not in ids  # filtre conservé
    rows = _csv_rows(admin_client.get(_query("csv")))
    bare = next(r for r in rows[1:] if r[0] == other.public_id)
    assert bare[-6:] == [NOT_PROVIDED] * 6


def test_followups_report_has_same_columns(admin_client, full_traveler, ebola_disease):
    QuarantineRecord.objects.create(
        traveler=full_traveler,
        disease=ebola_disease,
        started_on=ARRIVAL,
        expected_end_on=ARRIVAL + timedelta(days=21),
        address="Bouaké",
    )
    url = "/api/v1/reports/followups/?output=csv&date_from=2026-09-01&date_to=2026-10-31"
    header, row = _csv_rows(admin_client.get(url))
    assert header[-7:] == ITINERARY_HEADERS
    assert row[-7:] == [str(v) for v in itinerary_row(full_traveler)]
    assert admin_client.get(url.replace("csv", "xlsx")).status_code == 200


def test_export_permissions_are_unchanged(full_traveler, django_user_model, roles):
    user = django_user_model.objects.create_user(
        email="obs@example.ci",
        username="obs@example.ci",
        password="StrongPwd!2026",
    )
    RoleAssignment.objects.create(user=user, role=Role.objects.get(code=RoleCode.OBSERVER))
    c = APIClient()
    c.force_authenticate(user=user)
    for fmt in ("csv", "xlsx", "pdf"):
        assert c.get(_query(fmt)).status_code == 403
    assert APIClient().get(_query("xlsx")).status_code == 401


# ------------------------------------------------------- rapport hebdomadaire


def test_weekly_report_aggregates_without_personal_data(full_traveler):
    from apps.reports.services.weekly_aggregator import aggregate_weekly
    from apps.reports.services.weekly_excel import _render_csv_fallback, render_weekly_xlsx
    from apps.reports.services.weekly_pdf import render_weekly_pdf

    Traveler.objects.create(first_name="B", last_name="BARE")
    now = timezone.now()
    agg = aggregate_weekly(now - timedelta(days=7), now + timedelta(minutes=1))

    assert {"name": Country.objects.get(code="NG").name, "count": 1} in agg["by_origin_country"]
    assert {"name": NOT_PROVIDED, "count": 1} in agg["by_origin_country"]
    assert {"name": "Bouaké", "count": 1} in agg["by_destination_city"]
    assert agg["stay_duration"] == {
        "known": 1,
        "unknown": 1,
        "avg_days": 12.0,
        "min_days": 12,
        "max_days": 12,
    }

    wb = load_workbook(io.BytesIO(render_weekly_xlsx(agg)))
    assert {
        "Pays de provenance",
        "Pays de séjour récent",
        "Villes de destination",
        "Séjour en CI",
    } <= set(wb.sheetnames)
    stay = {r[0]: r[1] for r in wb["Séjour en CI"].iter_rows(min_row=2, values_only=True)}
    assert stay["Durée moyenne (jours)"] == 12.0
    csv_text = _render_csv_fallback(agg).decode("utf-8-sig")
    assert "=== Pays de provenance ===" in csv_text and "Durée moyenne (jours);12.0" in csv_text
    assert render_weekly_pdf(agg).startswith(b"%PDF")
    # Aucune donnée nominative dans les agrégats
    dumped = str(agg)
    assert full_traveler.public_id not in dumped and full_traveler.last_name not in dumped


def test_registration_date_uses_local_time(traveler):
    late = timezone.make_aware(datetime(2026, 10, 1, 23, 30), timezone.get_current_timezone())
    Traveler.objects.filter(pk=traveler.pk).update(created_at=late)
    traveler.refresh_from_db()
    assert itinerary_row(traveler)[0] == "01/10/2026"


# ------------------------------------------------- date de départ prévue (formulaire)


def test_declared_departure_drives_duration_and_wins_over_itinerary(full_traveler):
    full_traveler.planned_departure_date = ARRIVAL + timedelta(days=9)
    full_traveler.save()
    assert itinerary_row(full_traveler)[-1] == 9  # et non 12 (étape d'itinéraire)


def test_declared_departure_without_arrival_is_not_provided(traveler):
    traveler.arrival_date = None
    traveler.planned_departure_date = ARRIVAL
    traveler.save()
    assert itinerary_row(traveler)[-1] == NOT_PROVIDED


def _payload(**voyage):
    return {
        "voyage": {"arrival_date": "2026-10-01", "entry_point_code": "ABJ", **voyage},
        "identite": {"last_name": "DIOP", "first_name": "Awa", "gender": "F"},
        "historique": [{"role": "origin", "country_code": "FR", "city": "Paris"}],
        "confinement": {"city": "Abidjan"},
        "exposure": {},
        "symptoms": {},
        "declaration": {"truthful_declaration": True},
    }


def test_public_form_stores_planned_departure(pass_keys, entry_point):
    resp = APIClient().post(
        "/api/v1/ebola/public/register/",
        _payload(planned_departure_date="2026-10-15"),
        format="json",
    )
    assert resp.status_code == 201, resp.data
    t = Traveler.objects.get(public_id=resp.data["traveler"]["public_id"])
    assert t.planned_departure_date == date(2026, 10, 15)
    assert itinerary_row(t)[-1] == 14


@pytest.mark.parametrize("value", ["2026-09-30", "15/10/2026"])
def test_public_form_rejects_incoherent_or_invalid_departure(pass_keys, entry_point, value):
    resp = APIClient().post(
        "/api/v1/ebola/public/register/",
        _payload(planned_departure_date=value),
        format="json",
    )
    assert resp.status_code == 400
    assert not Traveler.objects.filter(last_name="DIOP").exists()


def test_public_form_without_departure_is_unchanged(pass_keys, entry_point):
    resp = APIClient().post("/api/v1/ebola/public/register/", _payload(), format="json")
    assert resp.status_code == 201, resp.data
    t = Traveler.objects.get(public_id=resp.data["traveler"]["public_id"])
    assert t.planned_departure_date is None
    assert itinerary_row(t)[-1] == NOT_PROVIDED


def test_admin_api_validates_planned_departure(traveler):
    from apps.travelers.serializers import TravelerSerializer

    traveler.arrival_date = ARRIVAL
    traveler.save()
    bad = TravelerSerializer(traveler, data={"planned_departure_date": "2026-09-01"}, partial=True)
    assert not bad.is_valid() and "planned_departure_date" in bad.errors
    ok = TravelerSerializer(traveler, data={"planned_departure_date": "2026-10-05"}, partial=True)
    assert ok.is_valid(), ok.errors
